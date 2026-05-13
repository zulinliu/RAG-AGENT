from __future__ import annotations

import os

from pptx import Presentation

from rag_qa.pipeline.parser_base import BaseParser, ParsedDocument, ParsedSection
from rag_qa.schemas.metadata import ChunkType


class PPTParser(BaseParser):

    async def parse(self, file_path: str) -> ParsedDocument:
        file_size, mime_type = self._get_file_info(file_path)
        checksum = self._compute_checksum(file_path)
        prs = Presentation(file_path)

        title = os.path.splitext(os.path.basename(file_path))[0]
        author = prs.core_properties.author or None
        if prs.core_properties.title:
            title = prs.core_properties.title

        sections: list[ParsedSection] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_title = f"幻灯片 {slide_idx + 1}"
            text_parts: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            text_parts.append(text)

                if shape.has_table:
                    table = shape.table
                    md_lines: list[str] = []
                    for row_idx, row in enumerate(table.rows):
                        cells = [
                            cell.text.strip().replace("\n", " ")
                            for cell in row.cells
                        ]
                        md_lines.append("| " + " | ".join(cells) + " |")
                        if row_idx == 0:
                            md_lines.append(
                                "| " + " | ".join(["---"] * len(cells)) + " |"
                            )
                    sections.append(
                        ParsedSection(
                            level=2,
                            title=f"{slide_title} - 表格",
                            content="\n".join(md_lines),
                            chunk_type=ChunkType.TABLE,
                        )
                    )

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    text_parts.append(f"[备注] {notes}")

            content = "\n".join(text_parts)
            if content.strip():
                sections.append(
                    ParsedSection(
                        level=1,
                        title=slide_title,
                        content=content,
                        chunk_type=ChunkType.PARAGRAPH,
                    )
                )

        return ParsedDocument(
            title=title,
            author=author,
            sections=sections,
            metadata={"slide_count": len(prs.slides)},
            file_path=file_path,
            file_size=file_size,
            mime_type=mime_type,
            checksum=checksum,
        )
