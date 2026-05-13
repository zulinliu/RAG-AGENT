from __future__ import annotations

import os

from docx import Document

from rag_qa.pipeline.parser_base import BaseParser, ParsedDocument, ParsedSection
from rag_qa.schemas.metadata import ChunkType


class WordParser(BaseParser):

    HEADING_LEVELS = {
        "Heading 1": 1,
        "Heading 2": 2,
        "Heading 3": 3,
        "Heading 4": 4,
        "Heading 5": 5,
        "Heading 6": 6,
        "标题 1": 1,
        "标题 2": 2,
        "标题 3": 3,
        "标题 4": 4,
        "标题 5": 5,
        "标题 6": 6,
    }

    async def parse(self, file_path: str) -> ParsedDocument:
        file_size, mime_type = self._get_file_info(file_path)
        checksum = self._compute_checksum(file_path)
        doc = Document(file_path)

        title = os.path.splitext(os.path.basename(file_path))[0]
        author = doc.core_properties.author or None
        if doc.core_properties.title:
            title = doc.core_properties.title

        sections: list[ParsedSection] = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                para = None
                for p in doc.paragraphs:
                    if p._element is element:
                        para = p
                        break
                if para is None:
                    continue

                style_name = para.style.name if para.style else ""
                text = para.text.strip()
                if not text:
                    continue

                level = self.HEADING_LEVELS.get(style_name)
                if level is not None:
                    sections.append(
                        ParsedSection(
                            level=level,
                            title=text,
                            content="",
                            chunk_type=ChunkType.HEADING,
                        )
                    )
                else:
                    sections.append(
                        ParsedSection(
                            level=0,
                            title="",
                            content=text,
                            chunk_type=ChunkType.PARAGRAPH,
                        )
                    )

            elif tag == "tbl":
                table = None
                for t in doc.tables:
                    if t._element is element:
                        table = t
                        break
                if table is None:
                    continue

                md_lines: list[str] = []
                for row_idx, row in enumerate(table.rows):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    md_lines.append("| " + " | ".join(cells) + " |")
                    if row_idx == 0:
                        md_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")

                sections.append(
                    ParsedSection(
                        level=0,
                        title="表格",
                        content="\n".join(md_lines),
                        chunk_type=ChunkType.TABLE,
                    )
                )

        return ParsedDocument(
            title=title,
            author=author,
            sections=sections,
            metadata={
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
            },
            file_path=file_path,
            file_size=file_size,
            mime_type=mime_type,
            checksum=checksum,
        )
