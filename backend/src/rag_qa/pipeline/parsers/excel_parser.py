from __future__ import annotations

import os

from openpyxl import load_workbook

from rag_qa.pipeline.parser_base import BaseParser, ParsedDocument, ParsedSection
from rag_qa.schemas.metadata import ChunkType


class ExcelParser(BaseParser):

    async def parse(self, file_path: str) -> ParsedDocument:
        file_size, mime_type = self._get_file_info(file_path)
        checksum = self._compute_checksum(file_path)
        wb = load_workbook(file_path, read_only=True, data_only=True)

        title = os.path.splitext(os.path.basename(file_path))[0]
        sections: list[ParsedSection] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            md_lines: list[str] = []
            header_written = False

            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                cells = []
                for cell in row:
                    if cell is None:
                        cells.append("")
                    else:
                        cells.append(str(cell).strip())
                md_lines.append("| " + " | ".join(cells) + " |")
                if not header_written:
                    md_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
                    header_written = True

            content = "\n".join(md_lines)
            if content.strip():
                sections.append(
                    ParsedSection(
                        level=1,
                        title=sheet_name,
                        content=content,
                        chunk_type=ChunkType.TABLE,
                    )
                )

        sheet_names = list(wb.sheetnames)
        wb.close()

        return ParsedDocument(
            title=title,
            author=None,
            sections=sections,
            metadata={"sheet_count": len(sheet_names), "sheet_names": sheet_names},
            file_path=file_path,
            file_size=file_size,
            mime_type=mime_type,
            checksum=checksum,
        )
